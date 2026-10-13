"""Small helpers shared by the deck builders (python-pptx + matplotlib mathtext)."""
from pathlib import Path
import hashlib

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TMP = Path("/tmp/ppt_assets"); TMP.mkdir(exist_ok=True)

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x44, 0x44, 0x44)
ACC = RGBColor(0xB2, 0x3A, 0x2A)
GREEN = RGBColor(0x2A, 0x7F, 0x3A)
CODE_BG = RGBColor(0xF3, 0xF4, 0xF6)
LIGHT = RGBColor(0xEE, 0xF2, 0xF7)


class Deck:
    def __init__(self, footer_text: str):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.blank = self.prs.slide_layouts[6]
        self.footer_text = footer_text
        self.n = 0

    # ------------------------------------------------------------ slides
    def slide(self, step: str, title: str):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.22), Inches(12.3), Inches(0.9))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        if step:
            r = p.add_run(); r.text = step + "  "
            r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = ACC
        r = p.add_run(); r.text = title
        r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = NAVY
        ln = s.shapes.add_shape(1, Inches(0.5), Inches(1.0), Inches(12.3), Emu(18000))
        ln.fill.solid(); ln.fill.fore_color.rgb = NAVY; ln.line.fill.background()
        fb = s.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(12.3), Inches(0.3))
        r = fb.text_frame.paragraphs[0].add_run()
        r.text = f"{self.footer_text}   ·   {self.n}"
        r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        return s

    def section(self, kicker: str, title: str, subtitle: str = ""):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        bg = s.shapes.add_shape(1, 0, 0, self.prs.slide_width, self.prs.slide_height)
        bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
        tb = s.shapes.add_textbox(Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.0))
        r = tb.text_frame.paragraphs[0].add_run(); r.text = kicker
        r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = RGBColor(0xF0, 0xB4, 0xA8)
        tb = s.shapes.add_textbox(Inches(0.9), Inches(2.9), Inches(11.5), Inches(1.4))
        tf = tb.text_frame; tf.word_wrap = True
        r = tf.paragraphs[0].add_run(); r.text = title
        r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
        if subtitle:
            tb = s.shapes.add_textbox(Inches(0.9), Inches(4.2), Inches(11.5), Inches(1.5))
            tf = tb.text_frame; tf.word_wrap = True
            r = tf.paragraphs[0].add_run(); r.text = subtitle
            r.font.size = Pt(18); r.font.color.rgb = RGBColor(0xDD, 0xE6, 0xF0)
        return s

    # ------------------------------------------------------------ content
    @staticmethod
    def bullets(slide, items, left, top, width, height, size=15, color=GREY, space=5):
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame; tf.word_wrap = True
        first = True
        for it in items:
            lvl = 0
            bold = False
            if isinstance(it, tuple):
                lvl, it = it
            if it.startswith("**"):
                bold = True; it = it[2:]
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            r = p.add_run()
            r.text = ("• " if lvl == 0 else "– ") + it
            r.font.size = Pt(size - 2 * lvl); r.font.color.rgb = color; r.font.bold = bold
            p.space_after = Pt(space)
            if lvl:
                p.level = lvl
        return tb

    @staticmethod
    def text(slide, txt, left, top, width, height, size=14, color=GREY, bold=False, italic=False, align=None):
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame; tf.word_wrap = True
        first = True
        for line in txt.split("\n"):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            if align is not None:
                p.alignment = align
            r = p.add_run(); r.text = line
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.italic = italic
        return tb

    @staticmethod
    def code(slide, code, left, top, width, height, size=11):
        box = slide.shapes.add_shape(1, left, top, width, height)
        box.fill.solid(); box.fill.fore_color.rgb = CODE_BG
        box.line.color.rgb = RGBColor(0xD0, 0xD0, 0xD0)
        tf = box.text_frame; tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = Inches(0.15)
        tf.margin_top = tf.margin_bottom = Inches(0.1)
        first = True
        for line in code.rstrip("\n").split("\n"):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run(); r.text = line if line else " "
            r.font.size = Pt(size); r.font.name = "Courier New"
            r.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
        return box

    @staticmethod
    def callout(slide, txt, left, top, width, height, size=13, color=ACC, fill=None):
        if fill is not None:
            box = slide.shapes.add_shape(1, left, top, width, height)
            box.fill.solid(); box.fill.fore_color.rgb = fill
            box.line.fill.background()
            tf = box.text_frame
            tf.margin_left = tf.margin_right = Inches(0.15)
            tf.vertical_anchor = MSO_ANCHOR.TOP
        else:
            tb = slide.shapes.add_textbox(left, top, width, height)
            tf = tb.text_frame
        tf.word_wrap = True
        first = True
        for line in txt.split("\n"):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            r = p.add_run(); r.text = line
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = True
        return tf

    @staticmethod
    def eq_png(tex, fontsize=22, color="#1F3A5F"):
        h = hashlib.md5((tex + str(fontsize) + color).encode()).hexdigest()[:10]
        p = TMP / f"eq_{h}.png"
        if p.exists():
            return p
        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, f"${tex}$", fontsize=fontsize, color=color)
        fig.savefig(p, dpi=300, bbox_inches="tight", pad_inches=0.05, transparent=True)
        plt.close(fig)
        return p

    def eq(self, slide, tex, left, top, height=Inches(0.6), max_width=None, fontsize=22):
        p = self.eq_png(tex, fontsize)
        with Image.open(p) as im:
            w_px, h_px = im.size
        if max_width is None:
            max_width = self.prs.slide_width - left - Inches(0.4)
        h = height
        w = int(h * w_px / h_px)
        if w > max_width:
            w = int(max_width); h = int(w * h_px / w_px)
        return slide.shapes.add_picture(str(p), left, top, width=w, height=h)

    @staticmethod
    def pic(slide, path, left, top, width=None, height=None):
        kw = {}
        if width is not None: kw["width"] = width
        if height is not None: kw["height"] = height
        return slide.shapes.add_picture(str(path), left, top, **kw)

    @staticmethod
    def table(slide, rows, left, top, width, col_widths, size=12, header=True):
        nrow, ncol = len(rows), len(rows[0])
        shp = slide.shapes.add_table(nrow, ncol, left, top, width, Inches(0.3) * nrow)
        tbl = shp.table
        for j, w in enumerate(col_widths):
            tbl.columns[j].width = w
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = tbl.cell(i, j); c.text = str(val)
                for p in c.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(size)
                        r.font.bold = header and i == 0
                c.margin_top = c.margin_bottom = Inches(0.03)
        return tbl

    def save(self, path):
        path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)
        return path
